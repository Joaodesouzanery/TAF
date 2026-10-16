#!/usr/bin/env python3 -I
"""Importa a obra 342 (Estudo REV10) -> web/data/obra342.json + supabase/seed/342.sql.

Uso: python3 -I scripts/import_planilhas.py <Estudo.xlsx> [saida_dir]
"""
import json, sys, datetime as dt
from pathlib import Path
import openpyxl

ESTUDO = Path(sys.argv[1])
ROOT = Path(__file__).resolve().parent.parent
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT
CONTRATO_342 = 3257970.74
ISS_TRIB = 0.085  # ISS 2% + lucro presumido 6,5% (doc rev01, sec. 5)
NAT = ["MO", "MAT", "EQP", "TER", "ADM"]  # colunas K..O da aba 31



# --- Plano de contas (doc rev01, seção 12) ----------------------------------
PLANO = [
 ("3","Receita e deduções",1,None),("3.1.01","Medições de contratos (BM)",3,None),("3.1.02","Reajuste contratual",3,None),("3.1.03","Aditivos e serviços extras",3,None),("3.2","Deduções: ISS, PIS, COFINS, glosas",2,None),
 ("4.1","Mão de obra de campo",2,"MO"),("4.1.01","Salários de produção, horas extras e adicionais",3,"MO"),("4.1.02","Encargos sociais e provisões",3,"MO"),("4.1.03","Benefícios: transporte, alimentação, cesta",3,"MO"),
 ("4.2","Materiais",2,"MAT"),("4.2.01","Tubos e conexões (FoFo, PVC, PEAD, concreto)",3,"MAT"),("4.2.02","Válvulas, registros, ventosas e peças especiais",3,"MAT"),
 ("4.2.03","Cimentos, agregados, argamassas e alvenaria",3,"MAT"),("4.2.04","Concreto usinado, aço e formas",3,"MAT"),("4.2.05","Elétrica e automação",3,"MAT"),
 ("4.2.06","Impermeabilização e geossintéticos",3,"MAT"),("4.2.07","Equipamentos permanentes fornecidos (bombas, guindaste)",3,"MAT"),("4.2.08","Material de consumo e ferramentas",3,"MAT"),
 ("4.3","Máquinas e veículos de produção",2,"EQP"),("4.3.01","Locação de máquinas, com ou sem operador",3,"EQP"),("4.3.02","Combustíveis e lubrificantes",3,"EQP"),("4.3.03","Manutenção, peças e hora de equipamento próprio",3,"EQP"),
 ("4.4","Terceiros e empreiteiros",2,"TER"),("4.4.01","Empreitada de mão de obra",3,"TER"),("4.4.02","Serviços técnicos: topografia, sondagem, ensaios",3,"TER"),("4.4.03","Transporte, carga e bota-fora",3,"TER"),
 ("4.5","Custo indireto da obra",2,"ADM"),("4.5.01","Equipe técnica e administrativa local",3,"ADM"),("4.5.02","Canteiro: mobilização, contêineres, ligações",3,"ADM"),("4.5.03","Alojamento, alimentação e utilidades",3,"ADM"),
 ("4.5.04","Segurança do trabalho, EPI e exames",3,"ADM"),("4.5.05","Veículos de apoio, TI, telefonia e impressão",3,"ADM"),("4.5.06","Projetos executivos, ART, taxas e licenças",3,"ADM"),("4.5.07","Seguros e garantias da obra",3,"ADM"),
 ("4.6","Provisão para perda em contrato oneroso",2,None),("5","Administração central",1,None),("6","Resultado financeiro",1,None),("7","IRPJ e CSLL",1,None),
]
MAT_PAC = {"FOFO":"4.2.01","PVC":"4.2.01","TCONC":"4.2.01","VALV":"4.2.02","PARAF":"4.2.08","ELET":"4.2.05","ELETS":"4.2.05","CONC":"4.2.04","ACO":"4.2.04","FORMA":"4.2.04","IMPER":"4.2.06","BOMBA":"4.2.07","GUIND":"4.2.07","CIVIL":"4.2.03"}
def conta_padrao(nat, pacote):
    """A natureza vem do cadastro; a conta é derivada (nunca digitada). Espelha conta_padrao() do SQL e do app."""
    if nat == "MO": return "4.1.01"
    if nat == "EQP": return "4.3.01"
    if nat == "TER": return "4.4.01"
    if nat == "ADM": return {"CANT": "4.5.02", "PROJ": "4.5.06"}.get(pacote, "4.5.01")
    return MAT_PAC.get(pacote, "4.2.08")

def d(v):
    return v.date().isoformat() if isinstance(v, (dt.datetime, dt.date)) else None


def rows(ws, header_row, max_col):
    it = ws.iter_rows(values_only=True, max_col=max_col)
    data = list(it)
    hdr = data[header_row - 1]
    return [dict(zip(hdr, r)) for r in data[header_row:] if any(c is not None for c in r)]


wb = openpyxl.load_workbook(ESTUDO, read_only=True, data_only=True)

# --- EAP (aba 31) ---------------------------------------------------------
eap, pai_stack = [], {}
for r in rows(wb["31. COMPOSIÇÃO POR ITEM"], 3, 16):
    cod = r["Código"]
    if not cod or cod == "GERAL":
        continue
    nivel = int(r["Nível"])
    pai_stack[nivel] = cod
    nat = {n: round(r[c] or 0, 2) for n, c in zip(
        NAT, ["Mão de obra", "Material e equipamento fornecido", "Máquinas e veículos",
              "Serviços de terceiros", "Indireto da obra"])}
    eap.append({
        "codigo": cod, "descricao": (r["Descrição"] or "").strip(), "unidade": r["Und"],
        "qtd": r["Qtd"], "nivel": nivel, "pai": pai_stack.get(nivel - 1) if nivel > 1 else None,
        "tipo_linha": r["Tipo"].lower(),
        "preco_venda_total": round(r["Preço licitado"] or 0, 2),
        "custo_direto_composicao": round(r["Custo direto"] or 0, 2),
        "natureza": nat,
    })
por_cod = {e["codigo"]: e for e in eap}

# --- Pacotes (aba 11) ------------------------------------------------------
pacotes = []
for r in rows(wb["11. META - PACOTES"], 3, 18):
    if not r.get("Pacote") or r["Pacote"] == "TOTAL":
        continue
    pacotes.append({
        "codigo": r["Pacote"], "nome": r["O que fechar"], "tipo": r["Tipo"],
        "responsavel": r["Responsável"], "eventos": r["Eventos"],
        "precisa_em_obra": d(r["Precisa em obra"]), "antecedencia_dias": r["Antecedência (dias)"],
        "fechar_ate": d(r["Fechar até"]), "meta": round(r["Meta de custo"] or 0, 2),
        "teto": round(r["Teto (margem zero)"] or 0, 2), "sem_cotacao": r["Sem cotação"] or 0,
        "condicao_pagamento": r["Condição de pagamento meta"], "observacoes": r["Observações"],
    })

# --- Meta por item (aba 12) ------------------------------------------------
sem_meta, itens_meta = [], 0
for r in rows(wb["12. META - ITENS"], 3, 21):
    cod = r.get("Item")
    if not cod or cod not in por_cod:
        continue
    e = por_cod[cod]
    qtd = r["Qtd"] or e["qtd"] or 0
    meta_total = round(r["Custo adotado"] if r["Custo adotado"] is not None else r["Meta custo total"], 2)
    e.update({
        "pacote": r["Pacote"], "evento": r["Evento"], "tipo": r["Tipo"],
        "meta_unit": round(meta_total / qtd, 4) if qtd else 0,
        "meta_total": meta_total, "teto_total": round(r["Teto total (margem zero)"] or 0, 2),
        "status_cotacao": r["Status"], "fonte_meta": r["Fonte da meta"],
        "precisa_em_obra": d(r["Precisa em obra"]), "fechar_ate": d(r["Fechar até"]),
    })
    itens_meta += 1

items = [e for e in eap if e["tipo_linha"] == "item"]
for e in items:
    tot = sum(e["natureza"].values())
    if tot > 0:
        e["composicao"] = [{"nat": k, "share": round(v / tot, 4)} for k, v in e["natureza"].items() if v > 0]
    else:  # item sem abertura na aba 31: assume a natureza pelo tipo
        e["composicao"] = [{"nat": "MAT" if "material" in (e.get("tipo") or "") else "TER", "share": 1.0}]
for e in items:
    if "meta_total" not in e:
        sem_meta.append(e["codigo"])
    # teto = preço licitado x (1 - 8,5%) quando a planilha não trouxe
    e.setdefault("teto_total", round(e["preco_venda_total"] * (1 - ISS_TRIB), 2))


# --- Fase 8: banco de MO (aba 18) e HH por unidade (aba 24) -------------------------
banco_mo = []
for r in wb["18. BANCO MO"].iter_rows(min_row=4, values_only=True):
    if not r[0] or r[12] in (None, ""): continue
    try: mensal = float(r[12])
    except (TypeError, ValueError): continue
    banco_mo.append({"funcao": str(r[0]).strip(), "equipe": r[1], "custo_mensal": round(mensal, 2), "custo_hora": round(mensal / 220, 2)})
hh = {}
for r in wb["24. CPU SANEAGO - ITENS"].iter_rows(min_row=4, values_only=True):
    if r[0] and isinstance(r[14], (int, float)) and isinstance(r[3], (int, float)) and r[3] > 0:
        hh[str(r[0]).strip()] = round(r[14] / r[3], 4)
for e in items:
    if e["codigo"] in hh: e["hh_unit"] = hh[e["codigo"]]

# --- Conferências ------------------------------------------------------------
soma_venda = round(sum(e["preco_venda_total"] for e in items), 2)
meta_pac = round(sum(p["meta"] for p in pacotes), 2)
meta_itens = round(sum(e.get("meta_total", 0) for e in items), 2)
rel = {
    "itens": len(items), "grupos": len(eap) - len(items), "pacotes": len(pacotes),
    "soma_preco_licitado": soma_venda, "contrato": CONTRATO_342,
    "diferenca_contrato": round(soma_venda - CONTRATO_342, 2),
    "meta_total_itens": meta_itens, "meta_total_pacotes": meta_pac,
    "dif_meta_itens_x_pacotes": round(meta_itens - meta_pac, 2),
    "itens_sem_meta": sem_meta,
    "itens_a_cotar": sum(1 for e in items if e.get("status_cotacao") == "A cotar"),
    "meta_a_cotar": round(sum(e.get("meta_total", 0) for e in items if e.get("status_cotacao") == "A cotar"), 2),
    "fonte": ESTUDO.name,
    "aviso": "META DE COMPRA enviada está marcada SUBSTITUÍDO pela REV02; usada a aba 11/12 do Estudo REV10. Aguardando REV13/META REV02.",
}
obra = {
    "codigo": "342", "nome": "UTR ETA Santo Antônio do Descoberto", "contratante": "SANEAGO",
    "contrato": "30000196/2026", "municipio": "Santo Antônio do Descoberto - GO",
    "valor_contrato": CONTRATO_342, "regime": "Preço global por evento (Anexo VII)",
    "tributos_nota": ISS_TRIB, "inicio_canteiro": "2026-10-19",
    "linha_base": {"versao": 1, "origem": "Estudo REV10 + META (aba 11/12)", "congelada": False},
}
(OUT / "web/data").mkdir(parents=True, exist_ok=True)
(OUT / "web/data/obra342.json").write_text(
    json.dumps({"obra": obra, "pacotes": pacotes, "eap": eap, "relatorio": rel,
                "banco_mo": banco_mo, "plano": [{"codigo": c, "nome": n, "nivel": l, "natureza": t} for c, n, l, t in PLANO]},
               ensure_ascii=False, indent=1), encoding="utf-8")

# --- SQL seed -------------------------------------------------------------
def q(v):
    if v is None: return "null"
    if isinstance(v, (int, float)): return repr(v)
    return "'" + str(v).replace("'", "''") + "'"

L = ["-- gerado por scripts/import_planilhas.py (obra 342)", "begin;",
     f"insert into obra(codigo,nome,contratante,contrato,municipio,valor_contrato,regime,tributos_nota) values "
     f"({q(obra['codigo'])},{q(obra['nome'])},{q(obra['contratante'])},{q(obra['contrato'])},{q(obra['municipio'])},"
     f"{obra['valor_contrato']},{q(obra['regime'])},{obra['tributos_nota']});"]
for p in pacotes:
    L.append("insert into pacote(obra_id,codigo,nome,tipo,responsavel,precisa_em_obra,fechar_ate,meta,teto,condicao_pagamento,observacoes) "
             f"select id,{q(p['codigo'])},{q(p['nome'])},{q(p['tipo'])},{q(p['responsavel'])},{q(p['precisa_em_obra'])},{q(p['fechar_ate'])},"
             f"{p['meta']},{p['teto']},{q(p['condicao_pagamento'])},{q(p['observacoes'])} from obra where codigo='342';")
for e in eap:
    L.append("insert into item_eap(obra_id,codigo,descricao,unidade,qtd,nivel,pai_codigo,eh_item,preco_venda_total,meta_total,teto_total,fonte_meta,status_cotacao,evento,pacote_id) "
             f"select o.id,{q(e['codigo'])},{q(e['descricao'])},{q(e['unidade'])},{q(e['qtd'])},{e['nivel']},{q(e['pai'])},"
             f"{str(e['tipo_linha']=='item').lower()},{e['preco_venda_total']},{q(e.get('meta_total'))},{q(e.get('teto_total'))},"
             f"{q(e.get('fonte_meta'))},{q(e.get('status_cotacao'))},{q(e.get('evento'))},"
             f"(select p.id from pacote p where p.obra_id=o.id and p.codigo={q(e.get('pacote'))}) from obra o where o.codigo='342';")
L.append("commit;")
pl = ["-- plano de contas (doc rev01, seção 12)", "begin;"] + [
    f"insert into conta(codigo,nome,nivel,natureza) values ({q(c)},{q(n)},{l},{q(t)}) on conflict do nothing;" for c, n, l, t in PLANO] + ["commit;"]
(OUT / "supabase/seed/plano_contas.sql").write_text("\n".join(pl), encoding="utf-8")
(OUT / "supabase/seed").mkdir(parents=True, exist_ok=True)
(OUT / "supabase/seed/342.sql").write_text("\n".join(L), encoding="utf-8")
print(json.dumps(rel, ensure_ascii=False, indent=1))
