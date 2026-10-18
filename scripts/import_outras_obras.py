#!/usr/bin/env python3 -I
"""Gera web/data/obra343.json (EAP da Viabilidade CAESB, só preço licitado) e web/data/obra281.json (histórico de lançamentos do consórcio).
Uso: python3 -I scripts/import_outras_obras.py fontes/Viabilidade_CAESB_Sobradinho_REV00.xlsm fontes/TAF_GESTAO_E_OPERACAO_rev6.xlsx"""
import json, sys, datetime as dt
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
via, taf = Path(sys.argv[1]), Path(sys.argv[2])
d = lambda v: v.date().isoformat() if isinstance(v, (dt.datetime, dt.date)) else None
ISS = 0.085

# ---------- 343: planilha de preços (aba PLANILHA) ----------
wb = openpyxl.load_workbook(via, read_only=True, data_only=True)
eap, grupo, vistos, n_g = [], None, set(), 0
for r in wb["PLANILHA"].iter_rows(min_row=2, values_only=True):
    seq, cod, desc, un, qtd, vu, tot = (list(r) + [None] * 7)[:7]
    if desc is None: continue
    if cod is None and qtd is None:   # linha de título
        n_g += 1; grupo = f"G{n_g:02d}"
        eap.append({"codigo": grupo, "descricao": str(desc).strip(), "nivel": 1, "pai": None, "tipo_linha": "grupo", "preco_venda_total": 0})
        continue
    if cod is None or not isinstance(qtd, (int, float)): continue
    c = str(cod).strip()
    if c in vistos: c = f"{c}.{n_g}"
    vistos.add(c)
    total = round(float(tot) if isinstance(tot, (int, float)) else float(qtd) * float(vu or 0), 2)
    eap.append({"codigo": c, "descricao": str(desc).strip(), "unidade": un, "qtd": float(qtd), "nivel": 2, "pai": grupo, "tipo_linha": "item",
                "preco_venda_total": total, "meta_total": None, "teto_total": round(total * (1 - ISS), 2), "meta_unit": None,
                "status_cotacao": "Sem estudo de custo", "fonte_meta": "Aguardando estudo de custo da obra 343"})
itens = [e for e in eap if e["tipo_linha"] == "item"]
soma = round(sum(e["preco_venda_total"] for e in itens), 2)
obra343 = {"codigo": "343", "nome": "Subadutora de Água Tratada Sobradinho I 192 (SAT.SB1.192)", "contratante": "CAESB", "contrato": "10214/2026",
           "municipio": "Sobradinho - DF", "valor_contrato": 5226634.03, "regime": "Preço unitário", "tributos_nota": ISS}
(ROOT / "web/data/obra343.json").write_text(json.dumps({"obra": obra343, "pacotes": [], "eap": eap, "plano": [], "relatorio": {
    "itens": len(itens), "soma_preco_licitado": soma, "contrato": 5226634.03, "diferenca_contrato": round(soma - 5226634.03, 2),
    "aviso": "Só preços da planilha de licitação. Meta de custo, composição e teto reais dependem do estudo de custo da obra 343."}}, ensure_ascii=False, indent=1), encoding="utf-8")
print("343:", len(itens), "itens, soma", soma)

# ---------- 281: histórico de lançamentos (aba 08D) ----------
wb2 = openpyxl.load_workbook(taf, read_only=True, data_only=True)
CONTA = {"MO": "4.1.01", "MAT": "4.2.08", "EQP": "4.3.01", "TER": "4.4.01", "ADM": "4.5.01"}
hist = []
for r in wb2["08D. LANÇAMENTOS 281"].iter_rows(min_row=4, values_only=True):
    n, tipo, emi, ven, _m, forn, desc, doc, val, cc, grupo, nat, sit = (list(r) + [None] * 13)[:13]
    if not isinstance(n, (int, float)) or not isinstance(val, (int, float)): continue
    nat = nat if nat in CONTA else "ADM"
    hist.append({"n": int(n), "tipo": tipo, "emissao": d(emi), "vencimento": d(ven), "fornecedor": str(forn or "").strip(), "descricao": str(desc or "").strip(),
                 "documento": str(doc or "").strip(), "valor": round(float(val), 2), "central": cc, "grupo": grupo, "nat": nat, "conta": CONTA[nat], "situacao": sit})
desp = [h for h in hist if h["tipo"] == "DESPESA"]
item_sem = {"codigo": "S.1", "descricao": "Lançamentos históricos sem item de orçamento (fila 'sem item')", "unidade": "UN", "qtd": 1, "nivel": 2, "pai": "S", "tipo_linha": "item",
            "preco_venda_total": 17120015.74, "meta_total": None, "teto_total": None, "status_cotacao": "Sem estudo de custo", "fonte_meta": "Obra 281 ainda sem orçamento de custo"}
grupo_s = {"codigo": "S", "descricao": "Histórico importado", "nivel": 1, "pai": None, "tipo_linha": "grupo", "preco_venda_total": 0}
obra281 = {"codigo": "281", "nome": "Ampliação do SAA de Mozarlândia/GO, Lote I (consórcio TAF Tecnobombas)", "contratante": "SANEAGO", "contrato": "30000415/2025",
           "municipio": "Mozarlândia - GO", "valor_contrato": 17120015.74, "regime": "Eventograma (semi-integrada)", "tributos_nota": ISS}
(ROOT / "web/data/obra281.json").write_text(json.dumps({"obra": obra281, "pacotes": [], "eap": [grupo_s, item_sem], "plano": [], "historico": hist,
    "relatorio": {"itens": 1, "lancamentos": len(hist), "despesas": len(desp), "total_despesas": round(sum(h["valor"] for h in desp), 2),
                  "aviso": "Histórico até 30/04/2026 do consórcio, vindo da aba 08D do TAF rev6. Sem orçamento de custo: tudo entra no item S.1."}}, ensure_ascii=False, indent=1), encoding="utf-8")
print("281:", len(hist), "lançamentos,", len(desp), "despesas, total", round(sum(h["valor"] for h in desp), 2))
